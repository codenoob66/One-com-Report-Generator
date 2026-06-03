import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import os
from openpyxl import load_workbook, Workbook

EMAIL_SHEET = "Agent updates Global"
MESSAGING_SHEET = "Assignee activity Messaging"


class ReportGeneratorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Excel Report Generator")
        self.root.geometry("850x800")

        self.messaging_path = tk.StringVar()
        self.email_path = tk.StringVar()
        self.hours_per_day = tk.DoubleVar(value=8.0)
        self.default_days = tk.DoubleVar(value=18.0)
        self.prod_goal = tk.DoubleVar(value=9.6)

        self.raw_data = {}

        # --- Main Paned Window ---
        main_paned_window = ttk.PanedWindow(root, orient=tk.VERTICAL)
        main_paned_window.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        top_controls_frame = tk.Frame(main_paned_window)
        main_paned_window.add(top_controls_frame, weight=0)

        # --- Input bar ---
        input_frame = tk.Frame(top_controls_frame)
        input_frame.pack(pady=(10, 0))

        tk.Label(input_frame, text="Hours/Day:").grid(row=0, column=0, padx=5)
        e = tk.Entry(input_frame, textvariable=self.hours_per_day, width=8)
        e.grid(row=0, column=1, padx=5)
        e.bind("<KeyRelease>", self.recalculate)

        tk.Label(input_frame, text="Default Days:").grid(row=0, column=2, padx=5)
        e = tk.Entry(input_frame, textvariable=self.default_days, width=8)
        e.grid(row=0, column=3, padx=5)
        e.bind("<KeyRelease>", self.recalculate)

        tk.Label(input_frame, text="PROD Goal:").grid(row=0, column=4, padx=5)
        e = tk.Entry(input_frame, textvariable=self.prod_goal, width=8)
        e.grid(row=0, column=5, padx=5)
        e.bind("<KeyRelease>", self.recalculate)

        # --- Messaging upload ---
        tk.Label(top_controls_frame, text="Messaging File:").pack(pady=(10, 0))
        messaging_frame = tk.Frame(top_controls_frame)
        messaging_frame.pack(pady=5)
        tk.Entry(messaging_frame, textvariable=self.messaging_path, width=50).pack(side=tk.LEFT, padx=(0, 5))
        tk.Button(messaging_frame, text="Browse...",
                  command=lambda: self.browse_file("Messaging", self.messaging_path)).pack(side=tk.LEFT)

        # --- Email upload ---
        tk.Label(top_controls_frame, text="Email File:").pack(pady=(10, 0))
        email_frame = tk.Frame(top_controls_frame)
        email_frame.pack(pady=5)
        tk.Entry(email_frame, textvariable=self.email_path, width=50).pack(side=tk.LEFT, padx=(0, 5))
        tk.Button(email_frame, text="Browse...",
                  command=lambda: self.browse_file("Email", self.email_path)).pack(side=tk.LEFT)

        # --- Filter names ---
        tk.Label(top_controls_frame, text="Search (partial match, comma separated):").pack(pady=(10, 0))
        self.filter_var = tk.StringVar(value="rafael")
        tk.Entry(top_controls_frame, textvariable=self.filter_var, width=50).pack(pady=5)

        # Buttons
        btn_frame = tk.Frame(top_controls_frame)
        btn_frame.pack(pady=10)
        tk.Button(btn_frame, text="Generate Report",
                  command=self.generate_report, bg="#4CAF50", fg="white").pack(side=tk.LEFT, padx=5)
        tk.Button(btn_frame, text="Team Heralds",
                  command=self.team_heralds, bg="#2196F3", fg="white").pack(side=tk.LEFT, padx=5)
        tk.Button(btn_frame, text="Export to Excel",
                  command=self.export_excel, bg="#FF9800", fg="white").pack(side=tk.LEFT, padx=5)

        # --- Tables frame ---
        tables_frame = tk.Frame(main_paned_window)
        main_paned_window.add(tables_frame, weight=1)

        paned_window = ttk.PanedWindow(tables_frame, orient=tk.VERTICAL)
        paned_window.pack(fill=tk.BOTH, expand=True)

        top_pane = tk.Frame(paned_window)
        paned_window.add(top_pane, weight=1)

        # --- Raw data table ---
        tk.Label(top_pane, text="Raw Data", font=("", 10, "bold")).pack(anchor="w")

        raw_frame = tk.Frame(top_pane)
        raw_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 5))

        columns = ("name", "messaging", "emails", "total", "days")
        self.tree = ttk.Treeview(raw_frame, columns=columns, show="headings", height=15)
        self.tree.heading("name", text="Name")
        self.tree.heading("messaging", text="Messaging")
        self.tree.heading("emails", text="Public Replies")
        self.tree.heading("total", text="Total")
        self.tree.heading("days", text="Days Worked")
        self.tree.column("name", width=200)
        self.tree.column("messaging", width=90, anchor="center")
        self.tree.column("emails", width=90, anchor="center")
        self.tree.column("total", width=90, anchor="center")
        self.tree.column("days", width=90, anchor="center")

        scrollbar = ttk.Scrollbar(raw_frame, orient=tk.VERTICAL, command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        self.tree.bind("<Double-1>", self.edit_days)

        bottom_pane = tk.Frame(paned_window)
        paned_window.add(bottom_pane, weight=1)

        # --- Metrics table ---
        tk.Label(bottom_pane, text="Productivity Metrics", font=("", 10, "bold")).pack(anchor="w")

        metrics_frame = tk.Frame(bottom_pane)
        metrics_frame.pack(fill=tk.BOTH, expand=True)

        metrics_cols = ("name", "hours", "total", "goal", "missing", "actual")
        self.metrics_tree = ttk.Treeview(metrics_frame, columns=metrics_cols, show="headings", height=15)
        self.metrics_tree.heading("name", text="Name")
        self.metrics_tree.heading("hours", text="Hours Worked")
        self.metrics_tree.heading("total", text="Total")
        self.metrics_tree.heading("goal", text="Tickets Goal")
        self.metrics_tree.heading("missing", text="Missing")
        self.metrics_tree.heading("actual", text="Actual Productivity")
        self.metrics_tree.column("name", width=200)
        self.metrics_tree.column("hours", width=90, anchor="center")
        self.metrics_tree.column("total", width=90, anchor="center")
        self.metrics_tree.column("goal", width=90, anchor="center")
        self.metrics_tree.column("missing", width=90, anchor="center")
        self.metrics_tree.column("actual", width=120, anchor="center")

        metrics_scrollbar = ttk.Scrollbar(metrics_frame, orient=tk.VERTICAL, command=self.metrics_tree.yview)
        self.metrics_tree.configure(yscrollcommand=metrics_scrollbar.set)
        self.metrics_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        metrics_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        self.edit_entry = None

    def browse_file(self, title, path_var):
        filename = filedialog.askopenfilename(
            title=f"Select {title} File",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
        )
        if filename:
            path_var.set(filename)

    def read_excel_counts(self, filepath, sheet_name):
        wb = load_workbook(filepath, data_only=True)
        if sheet_name not in wb.sheetnames:
            messagebox.showerror("Error", f"Sheet '{sheet_name}' not found in {filepath}")
            wb.close()
            return {}
        ws = wb[sheet_name]
        data = {}
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=2, values_only=True):
            name, count = row
            if name and name != "Updater name":
                data[str(name).strip()] = float(count) if count else 0.0
        wb.close()
        return data

    def team_heralds(self):
        self.filter_var.set("rafael, nevea, meg, marktwin, Niñoel Dagwayan, Rey Vergel, marcelo, lovely, Candy, Nathalie Dandan")
        self.generate_report()

    def generate_report(self):
        messaging_file = self.messaging_path.get()
        email_file = self.email_path.get()

        if not messaging_file and not email_file:
            messagebox.showwarning("Warning", "Please select at least one file!")
            return

        email_counts = {}
        messaging_counts = {}

        if email_file:
            if not os.path.exists(email_file):
                messagebox.showerror("Error", "Email file does not exist!")
                return
            email_counts = self.read_excel_counts(email_file, EMAIL_SHEET)

        if messaging_file:
            if not os.path.exists(messaging_file):
                messagebox.showerror("Error", "Messaging file does not exist!")
                return
            messaging_counts = self.read_excel_counts(messaging_file, MESSAGING_SHEET)

        filter_raw = self.filter_var.get().strip()
        filter_terms = [t.strip().lower() for t in filter_raw.split(",") if t.strip()] if filter_raw else []

        for item in self.tree.get_children():
            self.tree.delete(item)
        for item in self.metrics_tree.get_children():
            self.metrics_tree.delete(item)

        all_names = sorted(set(list(email_counts.keys()) + list(messaging_counts.keys())))
        if filter_terms:
            all_names = [n for n in all_names if any(term in n.lower() for term in filter_terms)]

        self.raw_data = {}
        default_days = self.default_days.get()

        for name in all_names:
            m = int(messaging_counts.get(name, 0))
            e = int(email_counts.get(name, 0))
            total = m + e
            self.raw_data[name] = {
                "messaging": m,
                "emails": e,
                "total": total,
                "days": default_days
            }
            self.tree.insert("", tk.END, values=(name, m, e, total, default_days))

        self.populate_metrics()

    def edit_days(self, event):
        if self.edit_entry:
            self.edit_entry.destroy()
            self.edit_entry = None

        region = self.tree.identify_region(event.x, event.y)
        if region != "cell":
            return
        col = self.tree.identify_column(event.x)
        if col != "#5":
            return
        item = self.tree.identify_row(event.y)
        if not item:
            return

        x, y, width, height = self.tree.bbox(item, "days")
        if not x:
            return
        value = self.tree.set(item, "days")

        self.edit_entry = tk.Entry(self.tree, width=8, justify="center")
        self.edit_entry.place(x=x, y=y, width=width, height=height)
        self.edit_entry.insert(0, value)
        self.edit_entry.select_range(0, tk.END)
        self.edit_entry.focus()

        def save_edit(event=None):
            new_value = self.edit_entry.get().strip()
            self.edit_entry.destroy()
            self.edit_entry = None
            try:
                days = float(new_value)
                self.raw_data[self.tree.set(item, "name")]["days"] = days
                self.tree.set(item, "days", days)
                self.populate_metrics()
            except ValueError:
                pass

        self.edit_entry.bind("<Return>", save_edit)
        self.edit_entry.bind("<FocusOut>", save_edit)
        self.edit_entry.bind("<Escape>", lambda e: self.edit_entry.destroy())

    def populate_metrics(self):
        for item in self.metrics_tree.get_children():
            self.metrics_tree.delete(item)

        hpd = self.hours_per_day.get()
        goal = self.prod_goal.get()

        for name, data in self.raw_data.items():
            hours = data["days"] * hpd
            total = data["total"]
            tickets_goal = hours * goal
            missing = int(tickets_goal - total)
            actual = round(total / hours, 2) if hours > 0 else 0
            self.metrics_tree.insert("", tk.END, values=(
                name, round(hours, 1), total, int(tickets_goal), missing, actual
            ))

    def recalculate(self, event=None):
        if not self.raw_data:
            return
        try:
            default_days = self.default_days.get()
            for data in self.raw_data.values():
                data["days"] = default_days
            for item in self.tree.get_children():
                name = self.tree.set(item, "name")
                self.tree.set(item, "days", self.raw_data[name]["days"])
            self.populate_metrics()
        except (ValueError, ZeroDivisionError):
            pass

    def export_excel(self):
        if not self.raw_data:
            messagebox.showwarning("Warning", "No data to export! Generate a report first.")
            return

        filepath = filedialog.asksaveasfilename(
            title="Save Excel Report",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")]
        )
        if not filepath:
            return

        hpd = self.hours_per_day.get()
        goal = self.prod_goal.get()

        wb = Workbook()

        ws1 = wb.active
        ws1.title = "Raw Data"
        ws1.append(["Name", "Messaging", "Public Replies", "Total", "Days Worked"])
        for name, data in self.raw_data.items():
            ws1.append([name, data["messaging"], data["emails"], data["total"], data["days"]])

        ws2 = wb.create_sheet("Productivity Metrics")
        ws2.append(["Name", "Days Worked", "Hours Worked", "Total", "Tickets Goal", "Missing", "Actual Productivity", "Hours per Day", "PROD Goal"])
        for name, data in self.raw_data.items():
            hours = data["days"] * hpd
            total = data["total"]
            tickets_goal = hours * goal
            missing = int(tickets_goal - total)
            actual = round(total / hours, 2) if hours > 0 else 0
            ws2.append([name, data["days"], round(hours, 1), total, int(tickets_goal), missing, actual, hpd, goal])

        ws3 = wb.create_sheet("Summary")
        ws3.append(["Metric", "Value"])
        ws3.append(["Hours per Day", hpd])
        ws3.append(["PROD Goal", goal])
        ws3.append(["Default Days Worked", self.default_days.get()])
        totals = [d["total"] for d in self.raw_data.values()]
        if totals:
            ws3.append(["Team Total", sum(totals)])

        wb.save(filepath)
        messagebox.showinfo("Success", f"Report exported to:\n{filepath}")


if __name__ == "__main__":
    root = tk.Tk()
    app = ReportGeneratorApp(root)
    root.mainloop()
